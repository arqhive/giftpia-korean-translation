"""4단계 대사 전수 추출 → translation/evt_texts.json, translation/evt_texts.xlsx

항목 = (명령, 원문 바이트) 하나. 위치는 저장하지 않는다(주입 때 역어셈블러로 다시 찾아 원문 바이트로 맞춘다).
순서: main.dol 의 맵 목록 순서(0x8018C240) → 나머지 맵, 파일 안에서는 위치 순.
"""
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import check_text
import evt
from dol import Dol

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
DEBUG_MAPS = {'map_debug.evt', 'map_01.evt', 'map_02.evt'}
KIND = {0x66: 'dialog', 0x88: 'dialog', 0x77: 'dialog', 0x7F: 'dialog', 0x89: 'dialog', 0x6D: 'dialog', 0x7E: 'choice'}


def map_order():
    D = Dol(os.path.join(ROOT, 'extract', 'main.dol'))
    names = []; a = 0x8018C240
    while True:
        p = D.u32(a)
        if p in (0, 0xFFFFFFFF): break
        names.append(D.read(p, 32).split(b'\0')[0].decode()); a += 4
    allf = sorted(x for x in os.listdir(os.path.join(ROOT, 'extract', 'dat', 'evt')) if x.endswith('.evt'))
    return names + [x for x in allf if x not in names]


def main():
    src = os.path.join(ROOT, 'extract', 'dat', 'evt')
    items = {}; order = []
    for fn in map_order():
        d = open(os.path.join(src, fn), 'rb').read()
        for kind, p, o, v, live in evt.texts(d):
            op = d[p]
            raw = d[o:v - 1] if kind == 'inline' else d[v:evt.text(d, v) - 1]
            key = (op if kind == 'inline' else 'p%02X' % op, raw)
            it = items.get(key)
            if it is None:
                it = items[key] = {'op': '%02X' % op, 'ptr': kind == 'ptr', 'raw': raw, 'files': [], 'n': 0, 'live': False}
                order.append(key)
            it['n'] += 1; it['live'] |= live
            if fn not in it['files']: it['files'].append(fn)
    out = []; names = map_order()
    for i, key in enumerate(order, 1):
        it = items[key]
        jp = evt.text_decode(it['raw'])
        m = re.match(r'(?:\{[0-9A-F]{2}\})*?\{0C\}\{([0-9A-F]{2})\}', jp)  # 앞머리 0C xx(표시 코드; 3E=포클 나레이션 등)
        flags = []
        if check_text.is_ticker(jp): flags.append('ticker')
        if '{04}' in jp: flags.append('size')
        if not it['live']: flags.append('dead')
        if set(it['files']) <= DEBUG_MAPS: flags.append('debugmap')
        out.append({'id': 'E%05d' % i, 'op': it['op'], 'ptr': it['ptr'], 'speaker': m.group(1) if m else '',
                    'kind': 'ticker' if 'ticker' in flags else KIND.get(int(it['op'], 16), 'dialog'),
                    'flags': flags, 'count': it['n'],
                    'files': ['ALL'] if len(it['files']) == len(names) else it['files'], 'jp': jp, 'ko': ''})
    os.makedirs(os.path.join(ROOT, 'translation'), exist_ok=True)
    json.dump(out, open(os.path.join(ROOT, 'translation', 'evt_texts.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    write_xlsx(out, os.path.join(ROOT, 'translation', 'evt_texts.xlsx'))
    chars = sum(len(re.sub(r'\{[0-9A-F]{2}\}|\n', '', x['jp'])) for x in out)
    print('항목', len(out), '글자', chars,
          '살아있음', sum(1 for x in out if 'dead' not in x['flags']),
          '디버그맵만', sum(1 for x in out if 'debugmap' in x['flags']),
          '라디오', sum(1 for x in out if x['kind'] == 'ticker'),
          '선택지', sum(1 for x in out if x['kind'] == 'choice'))


def write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook(); ws = wb.active; ws.title = 'evt'
    head = ['id', 'speaker', 'kind', 'flags', 'count', 'jp', 'ko', 'note']
    ws.append(head)
    for r in rows:
        ws.append([r['id'], r['speaker'], r['kind'], ','.join(r['flags']), r['count'], r['jp'], r['ko'], ''])
    for c, w in zip('ABCDEFGH', (9, 8, 8, 14, 7, 60, 60, 30)):
        ws.column_dimensions[c].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row[5:7]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    wb.save(path)


if __name__ == '__main__':
    main()
